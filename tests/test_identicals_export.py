# tests/test_identicals_export.py
"""Non-primary identicals report 0 target hours in exports and the audit log.

The real target stays in units.target_dept_hours so it comes back if the
primary changes; only what leaves the DB applies the identicals rule.
"""

import sqlite3

import pytest
from openpyxl import Workbook, load_workbook

from automation.create_db import create_database
from automation.export_to_workbook import export_to_workbook
from data.db import close_db, get_db
from data.loader import load_units
from data.writer import save_unit

# contract CN-1: 1001 is primary (earliest due), 1002/1003 are identicals.
UNITS = [
    ("1001", "CN-1", "2026-11-02", 30.0),
    ("1002", "CN-1", "2026-11-09", 30.0),
    ("1003", "CN-1", None, 30.0),
    ("2001", "CN-2", "2026-11-05", 12.0),  # single unit: not an identical
    ("3001", "", "2026-11-06", 8.0),  # no contract: never an identical
]


@pytest.fixture
def db(tmp_path):
    path = str(tmp_path / "identicals.db")
    create_database(path)
    conn = sqlite3.connect(path)
    conn.executemany(
        "INSERT INTO units (com_number, top_level_number, detailing_due_date, "
        "department_hours, target_dept_hours, detailer) VALUES (?, ?, ?, ?, ?, '')",
        [(com, cn, due, hrs, hrs) for com, cn, due, hrs in UNITS],
    )
    conn.commit()
    conn.close()
    yield path
    close_db()


def _view_targets(path):
    conn = get_db(path)
    return dict(conn.execute("SELECT com_number, target_dept_hours FROM v_reporting_export"))


def _stored_targets(path):
    conn = get_db(path)
    return dict(conn.execute("SELECT com_number, target_dept_hours FROM units"))


def test_reporting_view_zeroes_non_primary_identicals(db):
    assert _view_targets(db) == {
        "1001": 30.0,
        "1002": 0.0,
        "1003": 0.0,
        "2001": 12.0,
        "3001": 8.0,
    }
    # The DB keeps the real value.
    assert _stored_targets(db)["1002"] == 30.0


def test_view_matches_in_app_identicals_rule(db):
    in_app = {u.com_number for u in load_units(db) if u.is_non_primary_identical}
    zeroed = {com for com, hrs in _view_targets(db).items() if hrs == 0.0}
    assert zeroed == in_app == {"1002", "1003"}


def test_primary_change_restores_sibling_target(db):
    conn = get_db(db)
    conn.execute("UPDATE units SET detailing_due_date = '2026-12-31' WHERE com_number = '1001'")
    conn.commit()
    targets = _view_targets(db)
    assert targets["1002"] == 30.0
    assert targets["1001"] == 0.0


def test_old_view_definition_is_upgraded(tmp_path):
    path = str(tmp_path / "old.db")
    create_database(path)
    conn = sqlite3.connect(path)
    conn.execute("DROP VIEW IF EXISTS v_reporting_export")
    conn.execute(
        "CREATE VIEW v_reporting_export AS SELECT com_number, "
        "COALESCE(target_dept_hours, department_hours) AS target_dept_hours FROM units"
    )
    conn.executemany(
        "INSERT INTO units (com_number, top_level_number, detailing_due_date, department_hours) "
        "VALUES (?, 'CN-9', ?, 20.0)",
        [("9001", "2026-11-01"), ("9002", "2026-11-08")],
    )
    conn.commit()
    conn.close()
    try:
        assert _view_targets(path) == {"9001": 20.0, "9002": 0.0}
    finally:
        close_db()


def test_workbook_export_zeroes_non_primary_identicals(db, tmp_path):
    book = tmp_path / "schedule.xlsx"
    wb = Workbook()
    wb.active.title = "Current List"
    wb.save(book)

    export_to_workbook(db, str(book))

    ws = load_workbook(book)["Current List"]
    exported = {ws[f"C{r}"].value: ws[f"V{r}"].value for r in range(2, ws.max_row + 1)}
    assert exported == {"1001": 30.0, "1002": 0.0, "1003": 0.0, "2001": 12.0, "3001": 8.0}


def test_saving_non_primary_identical_does_not_audit_target_hours(db):
    unit = next(u for u in load_units(db) if u.com_number == "1002")
    assert unit.is_non_primary_identical
    unit.target_department_hours = 5.0  # whatever the form holds, it isn't audited
    unit.notes = "checked"
    save_unit(db, unit)

    conn = get_db(db)
    fields = {
        row[0]
        for row in conn.execute("SELECT field_name FROM _audit_log WHERE com_number = '1002'")
    }
    assert "notes" in fields
    assert "target_dept_hours" not in fields
    assert _stored_targets(db)["1002"] == 30.0


def test_saving_primary_still_audits_target_hours(db):
    unit = next(u for u in load_units(db) if u.com_number == "1001")
    unit.iec_internal_hours = 5.0
    save_unit(db, unit)

    conn = get_db(db)
    rows = conn.execute(
        "SELECT old_value, new_value FROM _audit_log "
        "WHERE com_number = '1001' AND field_name = 'target_dept_hours'"
    ).fetchall()
    assert [tuple(r) for r in rows] == [("30", "25")]
