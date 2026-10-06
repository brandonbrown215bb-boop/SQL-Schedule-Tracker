"""Tests for ExportService reporting feed export."""

import csv
from pathlib import Path
import sqlite3
import openpyxl
import pytest

from services.export_service import ExportService


def test_export_reporting_feed(db_with_units, tmp_path):
    svc = ExportService()
    feed_dir = tmp_path / "feed_out"
    count, xlsx_path = svc.export_reporting_feed(db_with_units, str(feed_dir))

    assert count > 0
    assert Path(xlsx_path).exists()
    csv_path = Path(xlsx_path).with_suffix(".csv")
    assert csv_path.exists()

    # Verify CSV headers
    with open(csv_path, "r", encoding="utf-8") as f:
        reader = csv.reader(f)
        headers = next(reader)
        assert "com_number" in headers
        assert "detailer_display" in headers
        assert "percent_complete" in headers
        assert "target_dept_hours" in headers
        assert "actual_hours" in headers

    # Verify XLSX headers and rows
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    assert "ReportingData" in wb.sheetnames
    ws = wb["ReportingData"]
    first_row = [cell.value for cell in next(ws.iter_rows())]
    assert "com_number" in first_row
    assert "detailer_display" in first_row
    wb.close()
